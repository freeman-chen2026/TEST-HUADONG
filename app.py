# -*- coding: utf-8 -*-
"""
专用：备案表生成（支持总调模板 & 华东区内模板，行号写死版）
"""

import streamlit as st
import pandas as pd
import io
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.cell.cell import MergedCell
import re
import json
from datetime import datetime, timedelta
import sys
import os
from copy import copy as _copy_style

# ---------- 页面设置 ----------
st.set_page_config(page_title="备案表自动生成", layout="wide")
st.title("📋 备案表自动生成（支持总调 / 华东区内）")

# ================================================================
# 内置数据与工具函数
# ================================================================
BUILTIN_CREW_DATA = [
    ("庚凡", "139 2463 9747", "430104197901184015", "430104197901184015"),
    ("张永一 / Yongyi ZHANG", "139 0125 9544", "110102196605202336", "110102196605202336"),
    ("马洪双 / Hongshuang MA", "156 2062 3177", "210726197306234110", "210726197306234110"),
    ("梅峰 / Feng MEI", "135 0967 8127", "17205/1 FCL", "510107197911242636"),
    ("王斌 / Bin WANG", "139 2527 2867", "3340398", "610104197911058331"),
    ("王少雄 / Shaoxiong WANG", "186 8387 9841", "510105198609042555", "510105198609042555"),
    ("苗旺旺 / Wangwang MIAO", "138 1871 5251", "410781198608019797", "410781198608019797"),
    ("赵岩松 / Yansong ZHAO", "186 1161 8385", "410103197004017014", "410103197004017014"),
    ("Bruce Roderick, WAINES", "186 6532 9796", "3448726", "000336198206158001"),
    ("Oliver Viktor, RACZ", "186 1197 3165", "000336198206158001", "000336198206158001"),
    ("Yiftah RAUCH / Yiftah, RAUCH", "186 1045 0563", "000972198112152001", "000972198112152001"),
    ("尤欣 / Xin YOU", "139 1608 5072", "620102197604293015", "620102197604293015"),
    ("李亚民 / Yamin Li", "133 6632 0878", "350104197107184915", "350104197107184915"),
    ("赵镭 / Lei ZHAO", "138 0883 9660", "440301198204157271", "440301198204157271"),
    ("彭罡", "186 1263 1888", "3240393", "440111198403244812"),
    ("胡君量 / Wan Leung WU / Kwan Leung WU", "132 6695 8816", "3025478", "124133200"),
    ("吴鹏", "136 1110 5901", "130103197602102115", "130103197602102115"),
    ("刘汇川", "188 5827 2791", "330103199003191618", "330103199003191618"),
    ("于龙飞", "156 5288 0812", "ZN00915", "210103198808123928"),
    ("林帅", "138 1156 6711", "110227198601130015", "110227198601130015"),
    ("张佳妮 / Jiani ZHANG", "136 6169 9966", "10183", "31010619840115002X"),
    ("张欢乐 / Huanle ZHANG", "186 1652 1529", "330381198705292523", "330381198705292523"),
    ("昝昭君 / Zhaojun ZAN", "182 9570 0579", "14272419950203313X", "14272419950203313X"),
    ("孙赫 / He SUN", "186 0102 1216", "410102197702243012", "410102197702243012"),
    ("李晓龙 / Xiaolong LI", "138 2378 1747", "420104197906150015", "420104197906150015"),
    ("李卉妍 / Huiyan LI", "138 5142 0321", "210106199007051542", "210106199007051542"),
    ("熊立凌 / Liling Xiong", "135 3821 6276", "421003199203222631", "421003199203222631"),
    ("HEALY, Darran William / Darran William HEALY", "852 6891 2350", "3141079", ""),
    ("ROEDER, SIMONE ELKE / SIMONE ELKE ROEDER", "852 6263 4569", "4213276", ""),
    ("王凯珮 / HOI PUI, WONG", "852 6858 8410", "10068", "Z411582(2)"),
    ("马坚 / Jian MA", "189 1770 2918", "320103196607089512", "320103196607089512"),
    ("卢江 / Jiang LU", "158 0045 6521", "420521198809280021", "420521198809280021"),
    ("孟周聪 / Zhoucong Meng", "135 6434 5029", "310113198307243215", "310113198307243215"),
    ("张帆 / Fan ZHANG", "138 0135 1294", "211002197306050057", "211002197306050057"),
    ("魏思远 / Siyuan WEI", "133 2110 4588", "230102198911054315", "230102198911054315"),
    ("王晟磊 / Shenglei WANG", "150 2689 7493", "310109198409264012", "310109198409264012"),
    ("Nathon Andrew G, NORBERG / Nathon Andrew G NORBERG", "186 0019 4610", "3050802", "A04700868"),
    ("Keith Robert, SHERREN / Keith Robert SHERREN", "137 3540 9744", "2755204", "P024165FF"),
    ("Rodolfo, BONETTI / Rodolfo BONETTI", "132 6284 1083", "12660", ""),
    ("危慧 / Hui WEI", "152 1349 1328", "43072119941115468X", "43072119941115468X"),
    ("李辛欣 / Xinxin LI", "135 5008 8666", "4209424", "510105198605023015"),
    ("Herve Daniel, STAMM / Herve Daniel STAMM", "183 1709 0300", "2666622", "M578620(2)"),
    ("樊婉程 / Wancheng FAN", "186 2017 4817", "440106199608180326", "440106199608180326"),
    ("刘爽 / Shuang LIU", "138 0125 8789", "110108196308296450", "110108196308296450"),
    ("刘凯 / Kai LIU", "135 2157 9157", "2833670", "230103198103275511"),
    ("詹佩佩 / Peipei ZHAN", "137 1440 5925", "10283", "440301198809275123"),
    ("花佩 / Pei HUA", "186 2631 0634", "320281198911117761", "320281198911117761"),
    ("翁英 / YING WENG", "130 6785 2000", "330106198801182024", "330106198801182024"),
    ("王莹 / Ying WANG", "159 1009 9069", "370125199004215621", "370125199004215621"),
    ("程佳俊 / Jiajun CHENG", "134 8010 3029", "511202198208161358", "511202198208161358"),
    ("蔡雨桐 / Yu Tong CHOI", "852 6426 7445", "10269", "V146532(5)"),
    ("俞凯 / Kai YU", "130 0579 0326", "120110196912180351", "120110196912180351"),
    ("杨杰 / Jie YANG", "186 1694 8903", "310104198411304413", "310104198411304413"),
    ("徐卓 / Zhuo XU", "139 1028 5510", "110105198906307113", "110105198906307113"),
    ("丁燕栒 / Yanxun DING", "135 6035 3829", "ZN00499", "440510199107260826"),
    ("万虹波 / Hongbo WAN", "133 1297 9906", "36050219860709003X", "36050219860709003X"),
    ("王国勤 / Guoqin WANG", "138 1815 8715", "2589322", "340822197610240215"),
    ("李潇恩 / Xiaoen LI", "158 0599 1600", "ZN00468", "510802199512100046"),
    ("孙辉 / Hui Sun", "139 1626 9572", "310228197810012612", "310228197810012612"),
    ("范蕾蕾 / Leilei FAN", "182 1000 6866", "ZN00347", "37010319861027002X"),
    ("张贺新 / Hexin ZHANG", "136 3773 1210", "420106197101020435", "420106197101020435"),
    ("李庆宏 / Qinghong LI", "135 0909 0503", "17260/1 FCL", "441402199107140256"),
    ("Eduard Pascal, Roski / Eduard Pascal Roski", "49 170 1534666", "3863535", "C9TR22VZM"),
    ("Peter Robert, JACKSON / Peter Robert JACKSON", "186 8245 1935", "000044197906253001", "000044197906253001"),
    ("廉卓群 / Zhuoqun LIAN", "133 5632 3949", "ZN00430", "370402199702018029"),
    ("孙浩 / Hao SUN", "136 7012 1990", "650103199102160037", "650103199102160037"),
    ("姚艳阁 / Yange YAO", "156 0127 9399", "10204", "130922199601151224"),
    ("茅邂文 / Xiewen MAO", "152 5181 7375", "ZN00903", "320683199911098627"),
    ("宋炜 / Wei SONG", "136 3256 5565", "37060219820621211X", "37060219820621211X"),
    ("BEEBE, Thaddeus John / Thaddeus John BEEBE", "852 6930 1609", "2743899", "R909452(A)"),
    ("张哲 / Zhe ZHANG", "139 0247 5026", "650104196604163310", "650104196604163310"),
    ("李海 / Hai LI", "136 8131 8388", "110105197201106130", "110105197201106130"),
    ("蔡国俊 / Kuo-Chun, TSAI", "86 157 1220 8304", "17203/1 FCL", "360019647"),
    ("杨涛 / Tao, YANG", "86 186 0122 5737", "140103197302010034", "140103197302010034"),
    ("朱正宇", "189 8335 3697", "350111197207152412", "350111197207152412"),
    ("金尚明 / Shangming JIN", "136 7113 8047", "210381197511034612", "210381197511034612"),
    ("赵婷婷", "138 2883 3162", "372524198212240023", "372524198212240023"),
    ("赖小燕 / Siau Mui LAI", "60 1239 05520", "A71366020", "A71366020"),
    ("何静文 / Ching Man, HO", "852 6421 0994", "Z630284(0)", "Z630284(0)"),
    ("周丽欢", "152 5710 6140", "330411199101195440", "330411199101195440"),
    ("AYA, MUGURUMA", "81 8071140700", "TT5868294", ""),
    ("梁广煜", "137 9428 7177", "EK9629672", ""),
    ("李园", "139 1178 3914", "110105198309149639", "110105198309149639"),
    ("孔铮", "139 1085 3981", "11010419801116125X", "11010419801116125X"),
    ("高峰", "135 8186 9017", "130826198001175332", "130826198001175332"),
    ("李阳", "133 6603 6567", "11010319850705091X", "11010319850705091X"),
    ("林生", "159 2162 9406", "350627198512262530", "350627198512262530"),
    ("谢依椿", "159 8942 4501", "441427198601221716", "441427198601221716"),
    ("廖关荣", "181 0755 9103", "360732199009095838", "360732199009095838"),
    ("林峰", "135 2260 7955", "150402198501122713", "150402198501122713"),
    ("丘东", "136 0044 6505", "11010119650406453x", "11010119650406453x"),
    ("王庆辉", "134 8013 9352", "350627198212052013", "350627198212052013"),
    ("姜磊", "135 1006 5318", "360502198312071333", "360502198312071333"),
    ("黄彦杰", "132 1157 2184", "450881199810196233", "450881199810196233"),
    ("王珍", "198 6662 9312", "622627199605283017", "622627199605283017"),
    ("赵国庆", "155 8849 2975", "370403200003133433", "370403200003133433"),
    ("王军", "853 62666900", "1458107(8)", "1458107(8)"),
    ("张德桃", "130 2883 6410", "360311199603192012", "360311199603192012"),
    ("江焰辉", "136 3148 0927", "440182199307270615", "440182199307270615"),
    ("孙龙", "156 9558 0691", "341623200010015613", "341623200010015613"),
    ("梁平", "138 2750 6225", "441223198510246217", "441223198510246217"),
    ("冯仁毫", "185 2028 6463", "440582199101153650", "440582199101153650"),
    ("焦石军", "139 2388 3525", "421224198310151013", "421224198310151013"),
    ("万子辰", "177 7005 7193", "", ""),
    ("卓辉", "157 7070 8632", "36073219981213009X", "36073219981213009X"),
    ("苏志斌", "159 0150 7150", "350782198308221539", "350782198308221539"),
    ("赵康", "191 6764 6172", "430321200303160170", "430321200303160170"),
    ("翟征宇", "134 1448 9793", "140211198612050031", "140211198612050031"),
    ("陈居瑜", "158 8962 6660", "460004197905030814", "460004197905030814"),
    ("林毅", "136 8642 0153", "350102197903213219", "350102197903213219"),
    ("Daniel, RICHTER", "374 9521 3832", "2556011", "C4W9RFH14"),
    ("郭春旭 / Guo Chunxu", "138 0136 1720", "110107197305150016", "110107197305150016"),
    ("黄海东", "138 0179 9315", "310105197506021215", "310105197506021215"),
]

CREW_COLUMNS = ["姓名", "联系方式", "执照号码", "证件号码"]
CREW_FILL_COLUMNS_GENERAL = ["职务", "姓名", "性别", "出生日期", "证件号码", "执照号码", "联系方式"]
CREW_FILL_COLUMNS_EAST = ["职务", "姓名", "性别", "出生日期", "国籍", "证件号码", "执照号码"]

NATION_MAP = {
    "CHN": "中国", "HKG": "中国香港", "TWN": "中国台湾省", "MAC": "中国澳门",
    "DEU": "德国", "USA": "美国", "GBR": "英国", "FRA": "法国", "RUS": "俄罗斯",
    "JPN": "日本", "KOR": "韩国", "SGP": "新加坡", "MYS": "马来西亚", "THA": "泰国",
    "VNM": "越南", "PHL": "菲律宾", "IDN": "印度尼西亚", "IND": "印度", "AUS": "澳大利亚",
    "CAN": "加拿大", "BRA": "巴西", "MEX": "墨西哥", "ZAF": "南非", "EGY": "埃及",
    "NGA": "尼日利亚", "KEN": "肯尼亚", "TZA": "坦桑尼亚", "ZWE": "津巴布韦",
    "NLD": "荷兰", "ITA": "意大利", "ESP": "西班牙", "PRT": "葡萄牙", "GRC": "希腊",
    "TUR": "土耳其", "SAU": "沙特阿拉伯", "ARE": "阿联酋", "ISR": "以色列",
    "IRN": "伊朗", "PAK": "巴基斯坦", "BGD": "孟加拉", "NPL": "尼泊尔",
    "LKA": "斯里兰卡", "MMR": "缅甸", "KHM": "柬埔寨", "LAO": "老挝", "MNG": "蒙古",
    "PRK": "朝鲜", "HUN": "匈牙利",
}

AIRCRAFT_TYPE_MAP = {
    "B3926": "LJ60", "B652R": "GLF4", "B8105": "GLEX", "B8160": "GLF5",
    "B8262": "GLF4", "B8292": "GLF5", "B8309": "GLF5", "MLLIN": "GLEX",
    "N2QE": "GL5T", "N328LM": "GL7T", "N550DR": "GLF5", "N577QT": "F900",
    "N7777U": "GLEX", "N777ZH": "GLF5", "N88AY": "GLF5", "T7178HT": "GL7T",
    "T7CJK": "GLEX", "VPCSZ": "GL7T", "VPCVA": "GLF6", "B652Q": "GLF4",
    "B652S": "GLF4", "B65AP": "GLF4",
}

def get_aircraft_type(reg, gd_type=""):
    reg_clean = str(reg).strip().upper() if reg else ""
    gd_type_str = str(gd_type).strip() if gd_type else ""
    if reg_clean and reg_clean in AIRCRAFT_TYPE_MAP:
        return AIRCRAFT_TYPE_MAP[reg_clean]
    return gd_type_str

def get_nation_name(code):
    code = str(code).strip().upper()
    return NATION_MAP.get(code, code)

def extract_chinese_name(full_name):
    if not full_name:
        return ""
    parts = full_name.split()
    chinese_parts = [p for p in parts if re.search(r'[\u4e00-\u9fff]', p)]
    if chinese_parts:
        return " ".join(chinese_parts)
    else:
        return full_name

def normalize_name(name):
    if not name:
        return ""
    name = re.sub(r'[\u4e00-\u9fff]+', '', name)
    name = re.sub(r'[,\s]+', ' ', name).strip().lower()
    return ' '.join(sorted(name.split()))

def _split_crew_name(name_val):
    return [p.strip() for p in re.split(r'\s*[/|、]\s*', name_val) if p.strip()]

def _find_crew_field(crew_name, field):
    if not crew_name:
        return ""
    target = str(crew_name).strip()
    target_cn = extract_chinese_name(target)
    target_norm = normalize_name(target)
    field_idx = CREW_COLUMNS.index(field)
    result = ""
    for row in BUILTIN_CREW_DATA:
        name_val = str(row[0] or "").strip()
        if not name_val:
            continue
        parts = _split_crew_name(name_val)
        matched = False
        if name_val == target:
            matched = True
        if not matched:
            for p in parts:
                if p == target:
                    matched = True; break
                if target_cn and p == target_cn:
                    matched = True; break
                if target_norm and normalize_name(p) == target_norm:
                    matched = True; break
        if matched:
            val = str(row[field_idx] or "").strip()
            if val:
                result = val
    return result

def find_contact(crew_name): return _find_crew_field(crew_name, "联系方式")
def find_license(crew_name): return _find_crew_field(crew_name, "执照号码")
def find_id(crew_name): return _find_crew_field(crew_name, "证件号码")

def is_18digit_id_card(val):
    if not val: return False
    s = re.sub(r'\s+', '', str(val))
    return bool(re.match(r'^[0-9]{17}[0-9Xx]$', s))

def parse_document_type(passport_no, doc_type):
    doc_type_str = str(doc_type).strip() if pd.notna(doc_type) else ""
    if doc_type_str:
        if "中华人民共和国居民身份证" in doc_type_str: return "身份证"
        if "港澳居民来往内地通行证" in doc_type_str: return "回乡证"
        return doc_type_str
    pn = str(passport_no).strip() if pd.notna(passport_no) else ""
    pn = re.sub(r'\s+', '', pn)
    if re.match(r'^[0-9]{15}$', pn) or re.match(r'^[0-9]{17}[0-9Xx]$', pn): return "身份证"
    return "护照"

# ★ 彻底修复 MergedCell 报错的终极安全函数
def safe_set_cell_value(ws, row, col, value):
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for merged_range in ws.merged_cells.ranges:
            if merged_range.min_row <= row <= merged_range.max_row and \
               merged_range.min_col <= col <= merged_range.max_col:
                target = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
                if not isinstance(target, MergedCell):
                    target.value = value
                return
    else:
        cell.value = value

def safe_set_cell_font(ws, row, col, font):
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        cell.font = font

def get_value_right(ws, row, start_col):
    for col in range(start_col, start_col + 10):
        cell = ws.cell(row=row, column=col)
        if cell.value and str(cell.value).strip():
            return str(cell.value).strip()
    return ""

def parse_utc_to_beijing(utc_str, date_str):
    try:
        time_part = utc_str.replace('Z', '').strip()
        if len(time_part) == 4:
            hour = int(time_part[:2]); minute = int(time_part[2:])
        elif len(time_part) == 3:
            hour = int(time_part[:1]); minute = int(time_part[1:])
        else: return "0000"
        day = int(re.search(r'\d+', date_str).group()) if re.search(r'\d+', date_str) else 1
        month_map = {"Jan":1,"Feb":2,"Mar":3,"Apr":4,"May":5,"Jun":6,"Jul":7,"Aug":8,"Sep":9,"Oct":10,"Nov":11,"Dec":12}
        month_str = re.search(r'[A-Za-z]{3}', date_str).group() if re.search(r'[A-Za-z]{3}', date_str) else "Jan"
        month = month_map.get(month_str[:3], 1)
        dt_beijing = datetime(2026, month, day, hour, minute) + timedelta(hours=8)
        return dt_beijing.strftime("%H%M")
    except: return "0000"

def parse_date_display(date_str):
    try:
        day = re.search(r'\d+', date_str).group()
        month_str = re.search(r'[A-Za-z]{3}', date_str).group()
        month_map = {"Jan":1,"Feb":2,"Mar":3,"Apr":4,"May":5,"Jun":6,"Jul":7,"Aug":8,"Sep":9,"Oct":10,"Nov":11,"Dec":12}
        month = month_map.get(month_str[:3], 1)
        return f"{month}月{int(day)}日"
    except: return date_str

def get_beijing_date_display(utc_time_str, date_str):
    if not utc_time_str or not date_str: return parse_date_display(date_str)
    try:
        time_part = utc_time_str.replace('Z', '').strip()
        if len(time_part) == 4:
            hour = int(time_part[:2]); minute = int(time_part[2:])
        elif len(time_part) == 3:
            hour = int(time_part[:1]); minute = int(time_part[1:])
        else: return parse_date_display(date_str)
        day = int(re.search(r'\d+', date_str).group())
        month_str = re.search(r'[A-Za-z]{3}', date_str).group()
        month_map = {"Jan":1,"Feb":2,"Mar":3,"Apr":4,"May":5,"Jun":6,"Jul":7,"Aug":8,"Sep":9,"Oct":10,"Nov":11,"Dec":12}
        month = month_map.get(month_str[:3], 1)
        dt_beijing = datetime(2026, month, day, hour, minute) + timedelta(hours=8)
        return f"{dt_beijing.month}月{dt_beijing.day}日"
    except: return parse_date_display(date_str)

def strip_single_letter_prefix(text):
    if text and re.match(r'^[A-Za-z]\s+', text): return re.sub(r'^[A-Za-z]\s+', '', text)
    return text

def parse_no_time_route(input_text, date_display):
    input_text = strip_single_letter_prefix(input_text)
    if not input_text or not input_text.strip(): return None
    text = input_text.strip()
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    if len(lines) >= 2:
        first_line = lines[0]; second_line = lines[1]
        flight_number = first_line.split()[0] if first_line.split() else None
        if not flight_number: return None
        parts = re.split(r'\s*[-—–]\s*', second_line)
        if len(parts) >= 2:
            dep_airport = parts[0].strip(); arr_airport = parts[1].strip()
            if dep_airport and arr_airport:
                return f"{date_display} {flight_number} {dep_airport}-{arr_airport}"
    else:
        flight_number = text.split()[0] if text.split() else None
        if not flight_number: return None
        time_pattern = (r'(?<![A-Za-z0-9])(\d{1,2}:\d{2}|\d{4})(?![A-Za-z0-9])'r'\s*[-—–]\s*'r'(\d{1,2}:\d{2}|\d{4})(?![A-Za-z0-9])')
        remaining = re.sub(time_pattern, '', text).strip()
        remaining = re.sub(r'\s*\+\s*\d+\s*', '', remaining).strip()
        airport_parts = re.split(r'\s*[-—–]\s*', remaining)
        if len(airport_parts) >= 2:
            dep_airport = airport_parts[-2].strip(); arr_airport = airport_parts[-1].strip()
            ch_dep = re.findall(r'[\u4e00-\u9fff]+', dep_airport)
            ch_arr = re.findall(r'[\u4e00-\u9fff]+', arr_airport)
            if ch_dep and ch_arr:
                dep_airport = ''.join(ch_dep); arr_airport = ''.join(ch_arr)
            if dep_airport and arr_airport:
                return f"{date_display} {flight_number} {dep_airport}-{arr_airport}"
    return None

def parse_with_time_route(input_text, date_display):
    input_text = strip_single_letter_prefix(input_text)
    if not input_text or not input_text.strip(): return input_text
    text = input_text.strip()
    time_pattern = (r'(?<![A-Za-z0-9])(\d{1,2}:\d{2}|\d{4})(?![A-Za-z0-9])'r'\s*[-—–]\s*'r'(\d{1,2}:\d{2}|\d{4})(?![A-Za-z0-9])')
    time_match = re.search(time_pattern, text)
    if time_match:
        dep_time = time_match.group(1).replace(':', ''); arr_time = time_match.group(2).replace(':', '')
        remaining = re.sub(time_pattern, '', text).strip()
    else:
        time_pattern2 = (r'(?<![A-Za-z0-9])(\d{1,2}:\d{2}|\d{4})(?![A-Za-z0-9])'r'\s+'r'(\d{1,2}:\d{2}|\d{4})(?![A-Za-z0-9])')
        time_match2 = re.search(time_pattern2, text)
        if time_match2:
            dep_time = time_match2.group(1).replace(':', ''); arr_time = time_match2.group(2).replace(':', '')
            remaining = re.sub(time_pattern2, '', text).strip()
        else: return input_text
    remaining = re.sub(r'\s*\+\s*\d+\s*', '', remaining).strip()
    airport_pattern = r'(.+?)\s*[-—–]\s*(.+)'
    airport_match = re.search(airport_pattern, remaining)
    if airport_match:
        dep_airport = airport_match.group(1).strip(); arr_airport = airport_match.group(2).strip()
        def extract_chinese(text):
            chinese = re.findall(r'[\u4e00-\u9fff]+', text)
            return ''.join(chinese) if chinese else text
        dep_airport = extract_chinese(dep_airport); arr_airport = extract_chinese(arr_airport)
        if dep_airport and arr_airport:
            return f"{date_display} {dep_airport}{dep_time}-{arr_time}{arr_airport}"
    return input_text

def parse_general_declaration(file_bytes):
    wb = load_workbook(file_bytes)
    ws = wb.active
    data = {}
    for row in ws.iter_rows(min_row=1, max_row=20):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                val = cell.value.strip()
                if "OPERATOR:" in val: data["operator"] = get_value_right(ws, cell.row, cell.column+1)
                elif "REG NO./FLT NO.:" in val:
                    reg_val = get_value_right(ws, cell.row, cell.column+1)
                    parts = reg_val.split()
                    data["reg"] = parts[0] if parts else reg_val
                    data["flt"] = parts[0] if parts else reg_val
                    if len(parts) > 1: data["flt"] = parts[1]
                elif "AC TYPE:" in val: data["ac_type_raw"] = get_value_right(ws, cell.row, cell.column+1)
                elif "FROM:" in val: data["from"] = get_value_right(ws, cell.row, cell.column+1)
                elif "TO:" in val: data["to"] = get_value_right(ws, cell.row, cell.column+1)
                elif "DATE/TIME:" in val:
                    date_time = get_value_right(ws, cell.row, cell.column+1)
                    data["date_time"] = date_time
                    if date_time:
                        parts = date_time.split()
                        data["utc_time"] = parts[0] if len(parts) > 0 else ""
                        data["date_str"] = parts[1] if len(parts) > 1 else ""
    data["ac_type"] = get_aircraft_type(data.get("reg", ""), data.get("ac_type_raw", ""))

    crew_data = []; passenger_data = []; section = None
    for row in ws.iter_rows(min_row=1):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                val = cell.value.strip()
                if "CREW MANIFEST" in val: section = 'crew'; break
                elif "PASSENGER MANIFEST" in val: section = 'passenger'; break
                elif "CARGO MANIFEST" in val or "DECLARATION OF HEALTH" in val: section = None; break
        if section == 'crew':
            first_cell = row[0]
            if first_cell.value and isinstance(first_cell.value, (int, float)) and len(row) >= 7:
                name_cell = row[1]
                if name_cell.value and isinstance(name_cell.value, str):
                    crew_data.append({"name": name_cell.value.strip(),"dob": row[2].value if row[2].value else "","gender": row[3].value if row[3].value else "","nationality": row[4].value if row[4].value else "","doc_type": row[5].value if row[5].value else "","passport_no": row[6].value if row[6].value else ""})
        elif section == 'passenger':
            first_cell = row[0]
            if first_cell.value and isinstance(first_cell.value, (int, float)) and len(row) >= 7:
                name_cell = row[1]
                if name_cell.value and isinstance(name_cell.value, str):
                    passenger_data.append({"name": name_cell.value.strip(),"dob": row[2].value if row[2].value else "","gender": row[3].value if row[3].value else "","nationality": row[4].value if row[4].value else "","doc_type": row[5].value if row[5].value else "","passport_no": row[6].value if row[6].value else ""})
    return data, crew_data, passenger_data

def _parse_hhmm(val):
    if val is None: return None
    try:
        if isinstance(val, float) and pd.isna(val): return None
    except: pass
    if hasattr(val, 'strftime'):
        try: return val.strftime('%H:%M')
        except: pass
    s = str(val).strip()
    m = re.match(r'^(\d{1,2}):(\d{2})', s)
    if m: return f"{int(m.group(1)):02d}:{m.group(2)}"
    return None

def parse_route_plan(file_bytes):
    try:
        df = pd.read_excel(file_bytes, skiprows=1)
        df.columns = [str(c).strip() for c in df.columns]
        if '飞机注册号' not in df.columns:
            df = pd.read_excel(file_bytes, header=1)
            df.columns = [str(c).strip() for c in df.columns]
        return df
    except Exception as e:
        st.warning(f"⚠️ 航段数据解析失败：{e}")
        return None

def _parse_gd_date(date_str):
    if not date_str: return None
    m = re.match(r'(\d{1,2})\s*([A-Za-z]{3})', str(date_str).strip())
    if not m: return None
    try: day = int(m.group(1))
    except: return None
    month_map = {"Jan":1,"Feb":2,"Mar":3,"Apr":4,"May":5,"Jun":6,"Jul":7,"Aug":8,"Sep":9,"Oct":10,"Nov":11,"Dec":12}
    month = month_map.get(m.group(2).capitalize()[:3])
    if not month: return None
    return (month, day)

def _extract_row_date(row_date):
    if row_date is None: return None
    try:
        if isinstance(row_date, float) and pd.isna(row_date): return None
    except: pass
    if hasattr(row_date, 'year'):
        try: return datetime(row_date.year, row_date.month, row_date.day)
        except: pass
    s = str(row_date).strip()
    m = re.match(r'(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})', s)
    if m:
        try: return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except: return None
    m = re.match(r'(\d{4})(\d{2})(\d{2})', s)
    if m:
        try: return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except: return None
    return None

def find_matching_flight(df, reg, from_airport, to_airport, gd_date=None):
    if df is None or df.empty: return None
    required = ['飞机注册号', '出发地', '到达地', '计划出发', '预计到达', '出发城市', '到达城市']
    if not all(c in df.columns for c in required): return None
    mask = ((df['飞机注册号'].astype(str).str.strip() == str(reg).strip()) &(df['出发地'].astype(str).str.strip().str.upper() == str(from_airport).strip().upper()) &(df['到达地'].astype(str).str.strip().str.upper() == str(to_airport).strip().upper()))
    matched = df[mask]
    if matched.empty: return None
    if gd_date is None or '出发日期' not in matched.columns: return matched.iloc[0]
    gd_month, gd_day = gd_date
    def _same_month_day(row_date):
        dt = _extract_row_date(row_date)
        return dt is not None and dt.month == gd_month and dt.day == gd_day
    exact = matched[matched['出发日期'].apply(_same_month_day)]
    if not exact.empty: return exact.iloc[0]
    def _nearby_month_day(row_date):
        dt = _extract_row_date(row_date)
        if dt is None: return False
        try: target = datetime(dt.year, gd_month, gd_day)
        except: return False
        return abs((dt - target).days) <= 1
    nearby = matched[matched['出发日期'].apply(_nearby_month_day)]
    if not nearby.empty: return nearby.iloc[0]
    return None

def build_route_from_flight(flight_row):
    reg = str(flight_row.get('飞机注册号', '') or '').strip()
    dep_time = _parse_hhmm(flight_row.get('计划出发'))
    arr_time = _parse_hhmm(flight_row.get('预计到达'))
    dep_city = str(flight_row.get('出发城市', '') or '').strip()
    arr_city = str(flight_row.get('到达城市', '') or '').strip()
    if not (reg and dep_time and arr_time and dep_city and arr_city): return None
    return f"F {reg} {dep_time.replace(':', '')} - {arr_time.replace(':', '')}  {dep_city} - {arr_city}"

def _style_name_cell(ws, row, col, text):
    target_row, target_col = row, col
    for merged_range in ws.merged_cells.ranges:
        if merged_range.min_row <= row <= merged_range.max_row and \
           merged_range.min_col <= col <= merged_range.max_col:
            target_row = merged_range.min_row
            target_col = merged_range.min_col
            break
    cell = ws.cell(row=target_row, column=target_col)
    if isinstance(cell, MergedCell): return
    text_str = str(text) if text is not None else ""
    new_align = _copy_style(cell.alignment)
    new_align.wrap_text = True
    cell.alignment = new_align
    n = len(text_str)
    if n <= 12: return
    if n <= 17: new_size = 10
    elif n <= 25: new_size = 9
    elif n <= 35: new_size = 8
    else: new_size = 7
    new_font = _copy_style(cell.font)
    new_font.size = new_size
    cell.font = new_font
    min_h = new_size * 2 + 4
    rd = ws.row_dimensions[target_row]
    current_h = rd.height
    if current_h is None or current_h < min_h: rd.height = min_h

def _style_snapshot(cell):
    return {'font': _copy_style(cell.font),'border': _copy_style(cell.border),'fill': _copy_style(cell.fill),'number_format': cell.number_format,'protection': _copy_style(cell.protection),'alignment': _copy_style(cell.alignment)}

def _style_apply(cell, snap):
    if isinstance(cell, MergedCell): return
    cell.font = _copy_style(snap['font'])
    cell.border = _copy_style(snap['border'])
    cell.fill = _copy_style(snap['fill'])
    cell.number_format = snap['number_format']
    cell.protection = _copy_style(snap['protection'])
    cell.alignment = _copy_style(snap['alignment'])

def _find_first_content_row(ws, start_row, max_scan=300):
    max_col = ws.max_column or 10
    for r in range(start_row, start_row + max_scan):
        for c in range(1, max_col + 1):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip(): return r
    return None

def _ensure_passenger_rows(ws, data_start_row, needed_count):
    if needed_count <= 0: return data_start_row
    from openpyxl.utils import get_column_letter
    end_row = _find_first_content_row(ws, data_start_row)
    if end_row is None: return data_start_row
    existing_rows = end_row - data_start_row
    if needed_count <= existing_rows: return data_start_row
    extra = needed_count - existing_rows
    template_row = end_row - 1
    max_col = ws.max_column or 10
    template_styles = {c: _style_snapshot(ws.cell(template_row, c)) for c in range(1, max_col + 1)}
    template_height = ws.row_dimensions[template_row].height
    max_row = ws.max_row
    saved_heights = {}
    for r in range(end_row, max_row + 1):
        h = ws.row_dimensions[r].height
        if h is not None: saved_heights[r] = h
    affected = []
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= end_row or mr.min_row < end_row <= mr.max_row:
            affected.append((mr.min_row, mr.min_col, mr.max_row, mr.max_col))
            try: ws.unmerge_cells(str(mr))
            except: pass
    ws.insert_rows(end_row, extra)
    
    for min_r, min_c, max_r, max_c in affected:
        if min_r >= end_row: new_min_r, new_max_r = min_r + extra, max_r + extra
        else: new_min_r, new_max_r = min_r, max_r + extra
        rng = (f"{get_column_letter(min_c)}{new_min_r}:" f"{get_column_letter(max_c)}{new_max_r}")
        try: ws.merge_cells(rng)
        except: pass
        
    for r in list(ws.row_dimensions.keys()):
        if r >= end_row:
            try: del ws.row_dimensions[r]
            except: pass
    for old_r, h in saved_heights.items(): ws.row_dimensions[old_r + extra].height = h
    
    template_merges = []
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row == template_row and mr.max_row == template_row:
            template_merges.append((mr.min_col, mr.max_col))

    for i in range(extra):
        new_row = end_row + i
        for c, snap in template_styles.items(): _style_apply(ws.cell(new_row, c), snap)
        if template_height: ws.row_dimensions[new_row].height = template_height
        for min_c, max_c in template_merges:
            rng = (f"{get_column_letter(min_c)}{new_row}:" f"{get_column_letter(max_c)}{new_row}")
            try: ws.merge_cells(rng)
            except: pass
            
    return data_start_row

# ================================================================
# 核心填充逻辑
# ================================================================
def fill_template_east_china(template_bytes, data, crew_rows, passenger_list, route_display):
    try:
        wb = load_workbook(template_bytes)
    except Exception as e:
        raise RuntimeError("❌ 模板文件解析失败。请确保你上传的是真正的 `.xlsx` 格式（请用 Excel 打开文件，选择“另存为”，格式选 Excel 工作簿(*.xlsx)），而不是直接修改后缀名。")

    ws = wb.active
    
    # 1. 飞行性质（第6行，B列）
    safe_set_cell_value(ws, 6, 2, "调机" if not passenger_list else "商务")
    
    # 2. 飞行信息（第10行，B=机型 C=注册号 D=航班号 E=行程）
    safe_set_cell_value(ws, 10, 2, data.get("ac_type", ""))
    safe_set_cell_value(ws, 10, 3, data.get("reg", ""))
    safe_set_cell_value(ws, 10, 4, data.get("flt", ""))
    safe_set_cell_value(ws, 10, 5, route_display if route_display else "")
    
    # 3. 机组信息（★ 华东模板固定从 15 行开始）
    role_map = {}
    for cr in crew_rows:
        role = str(cr.get("职务", "")).strip()
        if role and role not in role_map: role_map[role] = cr

    roles_pos = {"机长": 15, "副驾驶": 16, "乘务": 17, "机务": 18}
    for role_name, row_num in roles_pos.items():
        cr = role_map.get(role_name)
        if cr and str(cr.get("姓名", "")).strip() not in ("", "无"):
            name_val = str(cr.get("姓名", "")).strip()
            safe_set_cell_value(ws, row_num, 2, name_val)
            safe_set_cell_value(ws, row_num, 3, cr.get("性别", ""))
            safe_set_cell_value(ws, row_num, 4, cr.get("出生日期", ""))
            
            safe_set_cell_value(ws, row_num, 5, cr.get("国籍", ""))
            safe_set_cell_font(ws, row_num, 5, Font(name='宋体', size=16))
            
            safe_set_cell_value(ws, row_num, 6, cr.get("证件号码", ""))
            safe_set_cell_font(ws, row_num, 6, Font(name='宋体', size=16))
            
            safe_set_cell_value(ws, row_num, 7, cr.get("执照号码", ""))
            safe_set_cell_font(ws, row_num, 7, Font(name='宋体', size=12))
        else:
            for c in range(2, 8):
                safe_set_cell_value(ws, row_num, c, "无" if c == 2 else "")
    
    # 4. 乘客信息（★ 华东模板固定从 23 行开始）
    if passenger_list:
        start_row = 23
        passenger_start_row = _ensure_passenger_rows(ws, start_row, len(passenger_list))
        for i, pax in enumerate(passenger_list):
            row_num = passenger_start_row + i
            pax_name = extract_chinese_name(pax["name"])
            
            try:
                ws.merge_cells(start_row=row_num, start_column=6, end_row=row_num, end_column=7)
            except Exception:
                pass
                
            safe_set_cell_value(ws, row_num, 1, pax_name)
            safe_set_cell_value(ws, row_num, 2, pax.get("gender", ""))
            safe_set_cell_value(ws, row_num, 3, pax.get("dob", ""))
            safe_set_cell_value(ws, row_num, 4, get_nation_name(pax.get("nationality", "")))
            doc_type = pax.get("doc_type", "")
            doc_type_clean = parse_document_type("", doc_type) if (pd.notna(doc_type) and str(doc_type).strip()) else parse_document_type(pax.get("passport_no", ""), "")
            safe_set_cell_value(ws, row_num, 5, doc_type_clean)
            safe_set_cell_value(ws, row_num, 6, pax.get("passport_no", ""))
            
            for col_idx in range(1, 7):
                safe_set_cell_font(ws, row_num, col_idx, Font(name='宋体', size=16))

    output = BytesIO()
    wb.save(output); output.seek(0)
    return output

def fill_template_general(template_bytes, data, crew_rows, passenger_list, route_display):
    try:
        wb = load_workbook(template_bytes)
    except Exception as e:
        raise RuntimeError("❌ 模板文件解析失败。请确保你上传的是真正的 `.xlsx` 格式（请用 Excel 打开文件，选择“另存为”，格式选 Excel 工作簿(*.xlsx)），而不是直接修改后缀名。")

    ws = wb.active
    
    # 飞行目的
    if not passenger_list:
        for row in ws.iter_rows(min_row=1, max_row=10):
            for cell in row:
                if cell.value and isinstance(cell.value, str) and "飞行目的" in cell.value:
                    safe_set_cell_value(ws, cell.row + 1, 2, "调机")
                    break
    
    # 飞行信息
    info_row = None
    for row in ws.iter_rows(min_row=1, max_row=20):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.strip() in ["机型", "注册号", "航班号", "航班行程"]:
                info_row = cell.row; break
        if info_row: break
    
    if info_row:
        data_row = info_row + 1
        safe_set_cell_value(ws, data_row, 2, data.get("ac_type", ""))
        safe_set_cell_value(ws, data_row, 3, data.get("reg", ""))
        safe_set_cell_value(ws, data_row, 4, data.get("flt", ""))
        
        route_col = 5
        for cell in ws[info_row]:
            if cell.value and "行程" in str(cell.value):
                route_col = cell.column
                break
        safe_set_cell_value(ws, data_row, route_col, route_display if route_display else "")

    # 机组（★ 总调模板固定从 13 行开始）
    role_map = {}
    for cr in crew_rows:
        role = str(cr.get("职务", "")).strip()
        if role and role not in role_map: role_map[role] = cr

    roles_pos = {"机长": 13, "副驾驶": 14, "乘务": 15, "机务": 16}
    for role_name, row_num in roles_pos.items():
        cr = role_map.get(role_name)
        if cr and str(cr.get("姓名", "")).strip() not in ("", "无"):
            name_text = cr.get("姓名", "")
            safe_set_cell_value(ws, row_num, 2, name_text)
            safe_set_cell_value(ws, row_num, 3, cr.get("性别", ""))
            safe_set_cell_value(ws, row_num, 4, cr.get("出生日期", ""))
            safe_set_cell_value(ws, row_num, 5, cr.get("证件号码", ""))
            safe_set_cell_value(ws, row_num, 6, cr.get("执照号码", ""))
            safe_set_cell_value(ws, row_num, 7, cr.get("联系方式", ""))
            _style_name_cell(ws, row_num, 2, name_text)
        else:
            for c in range(2, 8): safe_set_cell_value(ws, row_num, c, "无" if c == 2 else "")

    # 乘客（★ 总调模板固定从 19 行开始）
    if passenger_list:
        start_row = 19
        passenger_start_row = _ensure_passenger_rows(ws, start_row, len(passenger_list))
        for i, pax in enumerate(passenger_list):
            row_num = passenger_start_row + i
            for col in range(1, 7): safe_set_cell_value(ws, row_num, col, None)
            pax_name = extract_chinese_name(pax["name"])
            safe_set_cell_value(ws, row_num, 1, pax_name)
            _style_name_cell(ws, row_num, 1, pax_name)
            safe_set_cell_value(ws, row_num, 2, pax.get("gender", ""))
            safe_set_cell_value(ws, row_num, 3, pax.get("dob", ""))
            safe_set_cell_value(ws, row_num, 4, get_nation_name(pax.get("nationality", "")))
            doc_type = pax.get("doc_type", "")
            doc_type_clean = parse_document_type("", doc_type) if (pd.notna(doc_type) and str(doc_type).strip()) else parse_document_type(pax.get("passport_no", ""), "")
            safe_set_cell_value(ws, row_num, 5, doc_type_clean)
            safe_set_cell_value(ws, row_num, 6, pax.get("passport_no", ""))

    output = BytesIO()
    wb.save(output); output.seek(0)
    return output

def fill_template(template_bytes, data, crew_rows, passenger_list, route_display, is_east_china):
    if is_east_china:
        return fill_template_east_china(template_bytes, data, crew_rows, passenger_list, route_display)
    else:
        return fill_template_general(template_bytes, data, crew_rows, passenger_list, route_display)

# ================================================================
# 页面逻辑
# ================================================================
data_file = st.file_uploader("上传 GD单（General Declaration）Excel（.xlsx）", type=["xlsx"], key="data")
template_file = st.file_uploader("上传模板（支持总调模板 或 华东区内模板，系统将自动识别）", type=["xlsx"], key="template")
flight_plan_file = st.file_uploader("（可选）上传航段数据导出（北京时间），用于预填文件名", type=["xlsx"], key="flight_plan")

if data_file and template_file:
    try:
        # 双重判断逻辑：优先看文件名，其次看表头
        is_east_china = False
        file_name = template_file.name.lower()
        
        if "华东" in file_name:
            is_east_china = True
        elif "总调" in file_name:
            is_east_china = False
        else:
            try:
                wb_temp = load_workbook(template_file)
                ws_temp = wb_temp.active
                for r in range(1, 30):
                    if ws_temp.cell(r, 1).value == "职务" and ws_temp.cell(r, 2).value == "姓名":
                        e_val = str(ws_temp.cell(r, 5).value or "")
                        if "国籍" in e_val:
                            is_east_china = True
                        else:
                            is_east_china = False
                        break
            except:
                pass

        if is_east_china:
            show_cols = CREW_FILL_COLUMNS_EAST
        else:
            show_cols = CREW_FILL_COLUMNS_GENERAL

        data, crew_list, passenger_list = parse_general_declaration(data_file)

        st.subheader("📋 本次机组信息（可编辑）")
        st.caption(
            "系统已从内置名单匹配出**将要写入的证件号码 / 执照号码 / 联系方式**，"
            "如有出入可直接在表格里修改（也可改姓名、性别、出生日期或调整职务）。"
            "改完下方下载按钮生成的就是最新数据。"
        )

        role_order = ["机长", "副驾驶", "乘务", "机务"]
        role_rows = {r: None for r in role_order}
        overflow_rows = []
        
        has_unknown_nation = False

        for idx, crew in enumerate(crew_list):
            name_cn = extract_chinese_name(crew["name"])
            id_val = find_id(crew["name"]) or crew.get("passport_no", "")
            lic_val = find_license(crew["name"])
            contact = find_contact(crew["name"])
            
            # ★ 新的国籍识别逻辑：优先使用 GD 单上的三字码，如果识别不了直接报错提示
            raw_nat = str(crew.get("nationality", "")).strip().upper()
            if raw_nat and raw_nat in NATION_MAP:
                nation = NATION_MAP[raw_nat]
            elif raw_nat:
                nation = f"待核实({raw_nat})"
                has_unknown_nation = True
            else:
                if is_18digit_id_card(id_val):
                    nation = "中国"
                else:
                    nation = "待核实"
                    has_unknown_nation = True

            if idx == 0: preferred_role = "机长"
            elif idx == 1: preferred_role = "副驾驶"
            else:
                gender = str(crew.get("gender", "")).strip()
                preferred_role = "乘务" if gender in ["女", "Female", "F"] else "机务"

            row = {
                "职务": preferred_role,
                "姓名": name_cn,
                "性别": crew.get("gender", ""),
                "出生日期": crew.get("dob", ""),
                "证件号码": id_val,
                "执照号码": lic_val,
                "联系方式": contact,
                "国籍": nation
            }

            if role_rows.get(preferred_role) is None: role_rows[preferred_role] = row
            else: overflow_rows.append(row)

        initial_rows = []
        for role in role_order:
            if role_rows[role] is not None: initial_rows.append(role_rows[role])
            else:
                empty_row = {"职务": role, "姓名": "无", "性别": "", "出生日期": "", "证件号码": "", "执照号码": "", "联系方式": "", "国籍": ""}
                initial_rows.append(empty_row)
        initial_rows.extend(overflow_rows)

        if overflow_rows:
            overflow_desc = "、".join(f"{r['姓名']}（当前标为「{r['职务']}」）" for r in overflow_rows)
            st.warning(f"⚠️ **本次机组共 {len(crew_list)} 人**，而模板机组区通常只有 4 行。以下人员**不会自动写入模板**：\n\n**{overflow_desc}**")

        # 弹出红色警告提醒用户去修改
        if has_unknown_nation and is_east_china:
            st.error("⚠️ **警告：部分机组人员的国籍未能自动识别，请在下方表格中手动修改！**")

        crew_signature = "|".join([c.get("name", "") for c in crew_list])
        editor_key = f"crew_fill_editor_{abs(hash(crew_signature)) % (10**8)}"

        df_fill = pd.DataFrame(initial_rows, columns=show_cols)
        
        col_config = {
            "职务": st.column_config.SelectboxColumn("职务", options=["机长", "副驾驶", "乘务", "机务"], width="small"),
            "姓名": st.column_config.TextColumn("姓名", width="medium"),
            "性别": st.column_config.TextColumn("性别", width="small"),
            "出生日期": st.column_config.TextColumn("出生日期", width="medium"),
            "证件号码": st.column_config.TextColumn("证件号码", width="medium"),
            "执照号码": st.column_config.TextColumn("执照号码", width="medium"),
        }
        if is_east_china:
            col_config["国籍"] = st.column_config.TextColumn("国籍", width="small")
        else:
            col_config["联系方式"] = st.column_config.TextColumn("联系方式", width="medium")

        edited_crew_df = st.data_editor(
            df_fill, 
            num_rows="dynamic", 
            width="stretch", 
            height=280, 
            key=editor_key,
            column_config=col_config
        )

        crew_for_template = edited_crew_df.fillna("").astype(str).to_dict("records") if not edited_crew_df.empty else []

        st.markdown("---")

        from_code = data.get("from", ""); to_code = data.get("to", "")
        date_str = data.get("date_str", ""); utc_time = data.get("utc_time", "")
        reg = data.get("reg", "")
        date_display = get_beijing_date_display(utc_time, date_str) if date_str else ""

        default_route = ""
        matched_note = ""
        gd_date_parsed = _parse_gd_date(date_str)
        if flight_plan_file is not None and reg and from_code and to_code:
            route_df = parse_route_plan(flight_plan_file)
            if route_df is not None:
                matched_flight = find_matching_flight(route_df, reg, from_code, to_code, gd_date=gd_date_parsed)
                if matched_flight is not None:
                    built = build_route_from_flight(matched_flight)
                    if built:
                        default_route = built
                        flight_no = str(matched_flight.get('航班号', '') or '').strip()
                        row_date = _extract_row_date(matched_flight.get('出发日期'))
                        date_hint = row_date.strftime("%m-%d") if row_date else "?"
                        matched_note = f"✅ 已从航段数据自动匹配到飞行计划（航班号 {flight_no}，{from_code} → {to_code}，{date_hint}），已预填到下方输入框"

        if not default_route:
            if date_str and from_code and to_code:
                bj_time = parse_utc_to_beijing(utc_time, date_str) if utc_time else "0000"
                default_route = f"{date_display} {from_code} {bj_time} XXXX {to_code}"
            else:
                default_route = f"{from_code}-{to_code}" if from_code and to_code else ""

        if not matched_note and flight_plan_file is not None and reg and from_code and to_code:
            date_hint = f"{gd_date_parsed[0]}月{gd_date_parsed[1]}日" if gd_date_parsed else "未知日期"
            st.info(f"ℹ️ 航段数据中未找到 {reg} {from_code}→{to_code}（{date_hint}）的匹配记录，已使用 GD单 信息生成默认值。")

        raw_route = st.text_input(
            "从Jetops复制航班信息并适当调整起落时间 比如： F B652S 08:00 - 14:00  柬埔寨金边 德崇 - 日本东京 羽田",
            value=default_route
        ).strip()

        with_time = parse_with_time_route(raw_route, date_display)
        route_display = with_time if (with_time != raw_route and with_time is not None) else (raw_route if raw_route else default_route)

        no_time = parse_no_time_route(raw_route, date_display)
        file_name_base = no_time if no_time is not None else (f"{date_display} {reg} {from_code}-{to_code}" if reg and from_code and to_code else route_display)

        safe_file_name = re.sub(r'[\\/*?:"<>|]', "_", file_name_base).strip() or "备案表"
        download_file_name = f"{safe_file_name}.xlsx"

        try:
            result_bytes = fill_template(template_file, data, crew_for_template, passenger_list, route_display, is_east_china).getvalue()
        except Exception as e:
            st.error(f"❌ 生成失败：{e}")
            st.stop()

        st.download_button(
            label="⬇️ 下载填充后的备案表",
            data=result_bytes,
            file_name=download_file_name,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        st.caption("⏳ 输入文件名之后稍等3秒钟后再点击下载")

    except Exception as e:
        st.error(f"❌ 处理失败：{e}")
        st.stop()
else:
    st.info("👆 请同时上传 GD单 和 模板文件。")
